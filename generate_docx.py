from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """
    Устанавливает границы ячейки, имитируя стиль Excel.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            for k, v in edge_data.items():
                element.set(qn('w:{}'.format(k)), str(v))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_color(cell, color_hex):
    """
    Устанавливает цвет фона ячейки.
    """
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_hex)
    shading_elm.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_font(cell, size=10, bold=False, color=None):
    """
    Устанавливает шрифт ячейки.
    """
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = 'Arial'
            if color:
                run.font.color.rgb = RGBColor(*color)

def create_excel_style_table(doc, data, headers, col_widths=None):
    """
    Создает таблицу в стиле Excel.
    data: список списков (данные)
    headers: список заголовков
    """
    table = doc.add_table(rows=1 + len(data), cols=len(headers))
    table.style = 'Table Grid'
    
    # Настройка заголовков
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_color(cell, 'D9D9D9') # Серый фон как в Excel
        set_cell_font(cell, size=10, bold=True)
    
    # Настройка данных
    for r, row_data in enumerate(data):
        for c, value in enumerate(row_data):
            cell = table.rows[r+1].cells[c]
            cell.text = str(value)
            set_cell_font(cell, size=10)
            
            # Пример: выделение цветом для конкретных ячеек (например, большие числа)
            if value == "12 млн":
                set_cell_color(cell, 'E2EFDA') # Светло-зеленый фон

    # Настройка ширины колонок
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)
    
    return table

# --- Создание документа ---
doc = Document()

# Заголовок
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(12)

title = doc.add_heading('Анализ системы АБС ЦФТ-Банк', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Раздел 1: Статистика
doc.add_heading('1. Масштаб системы', level=1)

stats_data = [
    ["Таблиц", "8 000"],
    ["Методов", "50 000"],
    ["Представлений", "13 000"],
    ["Строк кода дистрибутива", "12 млн"],
]
stats_headers = ["Параметр", "Значение"]
col_widths_stats = [10, 6]

create_excel_style_table(doc, stats_data, stats_headers, col_widths_stats)

# Раздел 2: Этапы миграции
doc.add_heading('2. Этапы проекта', level=1)

phases_data = [
    ["Этап 1", "Адаптация кода под DBI", "Анализ, регрессионные тесты, конвертация PL/SQL"],
    ["Этап 2", "Миграция данных и переход", "CutWYN SBS, Oracle -> PostgreSQL, промышленная платформа"],
    ["Этап 3", "Тестирование и валидация", "DBI Regress, DBI perform, проверка на целевых объемах"],
]
phases_headers = ["Этап", "Задача", "Детали"]
col_widths_phases = [2, 6, 12]

create_excel_style_table(doc, phases_data, phases_headers, col_widths_phases)

# Сохранение
output_filename = 'DBI_Analysis_Report.docx'
doc.save(output_filename)
print(f"Документ '{output_filename}' успешно создан.")
