#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import openpyxl
import os
from pathlib import Path

# Конфигурация
INPUT_FILE = "F:\TO_DBI\DATA\Patch_USERNAME\3_methods2refactor.xlsx"  # Замените на ваш файл
OUTPUT_DIR = "F:\TO_DBI\DATA\Patch_USERNAME\"  # Каталог для сохранения .pck файлов
SHEET_NAME = "Select methods"  # Имя листа с данными (если None, берётся активный)

def normalize_user(user):
    """Приводит имя пользователя к верхнему регистру и удаляет пробелы"""
    if user is None:
        return "UNKNOWN"
    return str(user).strip().upper()

def create_pck_file(output_path, class_id, short_name):
    """Создаёт .pck файл с заданной структурой"""
    content = f"""VER2
REM Список элементов
REM CFT-Platform-IDE: 2.36.406

METH {class_id} {short_name}
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(content)

def main():
    # Создаём выходной каталог
    Path(OUTPUT_DIR).mkdir(parents=True, exist_ok=True)
    
    # Открываем Excel файл
    print(f"Загрузка файла: {INPUT_FILE}")
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    
    # Выбираем лист
    if SHEET_NAME:
        sheet = wb[SHEET_NAME]
    else:
        sheet = wb.active
    
    # Получаем заголовки (первая строка)
    headers = []
    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=col).value
        headers.append(header)
    
    # Определяем индексы нужных колонок
    col_idx = {}
    for idx, header in enumerate(headers):
        if header == "Status X":
            col_idx['status_x'] = idx + 1
        elif header == "USER_MODIFIED":
            col_idx['user_modified'] = idx + 1
        elif header == "CLASS_ID":
            col_idx['class_id'] = idx + 1
        elif header == "SHORT_NAME":
            col_idx['short_name'] = idx + 1
    
    # Проверяем, что все колонки найдены
    required = ['status_x', 'user_modified', 'class_id', 'short_name']
    missing = [r for r in required if r not in col_idx]
    if missing:
        print(f"ОШИБКА: Не найдены колонки: {missing}")
        return
    
    print(f"Найдены колонки: {col_idx}")
    
    # Собираем данные со 2-й строки
    rows_data = []
    for row in range(2, sheet.max_row + 1):
        status_x = sheet.cell(row=row, column=col_idx['status_x']).value
        # Пропускаем строки, где Status X не пустой
        if status_x is not None and str(status_x).strip() != '':
            continue
        
        user_modified = sheet.cell(row=row, column=col_idx['user_modified']).value
        class_id = sheet.cell(row=row, column=col_idx['class_id']).value
        short_name = sheet.cell(row=row, column=col_idx['short_name']).value
        
        # Пропускаем строки с пустыми обязательными полями
        if class_id is None or short_name is None:
            continue
        
        rows_data.append({
            'user_modified': user_modified,
            'class_id': str(class_id).strip(),
            'short_name': str(short_name).strip()
        })
    
    # Сортировка: по USER_MODIFIED (без учёта регистра), затем по CLASS_ID, затем по SHORT_NAME
    rows_data.sort(key=lambda x: (
        x['user_modified'].lower() if x['user_modified'] else '',
        x['class_id'].lower(),
        x['short_name'].lower()
    ))
    
    # Группировка по USER_MODIFIED
    groups = {}
    for row in rows_data:
        user = normalize_user(row['user_modified'])
        if user not in groups:
            groups[user] = []
        groups[user].append(row)
    
    print(f"\nНайдено {len(rows_data)} записей для {len(groups)} пользователей")
    
    # Создание .pck файлов
    for user, items in groups.items():
        filename = f"patch_{user}.pck"
        filepath = os.path.join(OUTPUT_DIR, filename)
        
        # Формируем содержимое файла
        lines = []
        lines.append("VER2")
        lines.append("REM Список элементов")
        lines.append("REM CFT-Platform-IDE: 2.36.406")
        lines.append("")
        
        for item in items:
            lines.append(f"METH {item['class_id']} {item['short_name']}")
        
        # Добавляем пустую строку в конце
        lines.append("")
        
        # Записываем файл
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))
        
        print(f"  Создан: {filename} ({len(items)} записей)")
    
    print(f"\nГотово! Файлы сохранены в: {OUTPUT_DIR}")

if __name__ == "__main__":
    main()